"""DOCX document loader using python-docx.

Extracts paragraph text and table content.  DOCX files do not have
native page boundaries, so the entire document is represented as a
single logical page.
"""

from __future__ import annotations

import logging
from pathlib import Path

import docx

from rag.ingestion.document_id import generate_document_id
from rag.ingestion.exceptions import InvalidDocumentError
from rag.ingestion.models import Document, PageContent

logger = logging.getLogger(__name__)


def load_docx(file_path: Path, *, document_name: str | None = None) -> Document:
    """Parse a DOCX file and return a normalized ``Document``.

    Args:
        file_path: Path to a validated DOCX file.
        document_name: Display name for the document.  Falls back to
            the file's basename.

    Raises:
        InvalidDocumentError: If python-docx cannot open or read the file.
    """
    name = document_name or file_path.name

    try:
        doc = docx.Document(str(file_path))
    except Exception as exc:
        raise InvalidDocumentError(f"Failed to open DOCX '{name}': {exc}") from exc

    parts: list[str] = []

    # Extract paragraphs.
    try:
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                parts.append(text)
    except Exception as exc:
        raise InvalidDocumentError(
            f"Error extracting paragraphs from DOCX '{name}': {exc}"
        ) from exc

    # Extract table content.
    try:
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                row_text = " | ".join(cells)
                if row_text.strip():
                    parts.append(row_text)
    except Exception as exc:
        raise InvalidDocumentError(
            f"Error extracting tables from DOCX '{name}': {exc}"
        ) from exc

    content = "\n".join(parts)
    file_size = file_path.stat().st_size

    page = PageContent(
        page_number=1,
        content=content,
        metadata={"char_count": len(content)},
    )

    return Document(
        document_id=generate_document_id(),
        document_name=name,
        source_type="docx",
        pages=[page],
        metadata={
            "total_pages": 1,
            "file_size_bytes": file_size,
        },
    )
